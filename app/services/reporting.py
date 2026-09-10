"""Post-analysis reporting: summary PDF + 3-tab Excel, stored and (dry-run)
emailed to the uploading Manager. Executives/MA access the same artifacts
via run-scoped download endpoints."""

import io
from datetime import UTC, datetime

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.logging import get_logger
from app.models.analysis import (
    AnalysisRun,
    EmailStatus,
    GeneratedReport,
    ReportKind,
    ViolationResult,
)
from app.models.user import User
from app.services.storage import build_key, get_storage

log = get_logger("reporting")

TAB_SENT = "Sent"
TAB_NO_EMAIL = "Violations-No-Email"
TAB_SKIPPED = "Skipped-Not-Processed"


async def _run_rows(db: AsyncSession, run: AnalysisRun):
    stmt = (
        select(ViolationResult)
        .where(ViolationResult.run_id == run.id)
        .order_by(ViolationResult.employee_code)
    )
    return list((await db.execute(stmt)).scalars().all())


# --- 3-tab Excel ---------------------------------------------------------------


_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\x0b")


def _escape_formula(value) -> str:
    """Neutralize spreadsheet-formula injection in cells written to Excel."""
    if not isinstance(value, str):
        value = str(value)
    if value.lstrip(" ")[:1] in _FORMULA_PREFIXES:
        return "'" + value
    return value


def build_excel_3tab(run: AnalysisRun, violations, skips: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = TAB_SENT
    header_font = Font(bold=True)

    def write_sheet(ws, headers, rows):
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
        for r in rows:
            ws.append([_escape_formula(c) for c in r])

    sent = [
        [v.employee_code, v.employee_name or "", v.email_to or "",
         v.violation_reason, v.sent_at.isoformat() if v.sent_at else ""]
        for v in violations if v.email_status == EmailStatus.SENT
    ]
    write_sheet(
        ws,
        ["Employee ID", "Name", "Emailed To", "Reason", "Sent At"],
        sent,
    )

    no_email = [
        [v.employee_code, v.employee_name or "", v.violation_reason]
        for v in violations if v.email_status == EmailStatus.SKIPPED_NO_EMAIL
    ]
    write_sheet(wb.create_sheet(TAB_NO_EMAIL), ["Employee ID", "Name", "Reason"], no_email)

    skipped = [
        [s.get("employee_code") or "(unreadable)", s.get("status", ""), s.get("reason", "")]
        for s in skips
    ]
    write_sheet(
        wb.create_sheet(TAB_SKIPPED),
        ["Employee ID", "Skip Reason", "Detail"],
        skipped,
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_summary_pdf(run: AnalysisRun, violations) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    totals = [
        ("Files processed", run.files_processed),
        ("Rows processed", run.rows_processed),
        ("Matched to repository", run.matched),
        ("Violations detected", run.violations),
        ("Emails sent", run.emails_sent),
        ("Violators without email", run.emails_missing),
        ("Email errors", run.emails_errored),
        ("Skipped / not processed", run.skipped),
    ]

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    c.setTitle("VeriTrack AI — GETS Analysis Summary")
    y = height - 30 * mm

    c.setFont("Helvetica-Bold", 16)
    c.drawString(25 * mm, y, "VeriTrack AI — GETS Analysis Summary")
    y -= 10 * mm
    c.setFont("Helvetica", 10)
    c.drawString(25 * mm, y, f"Run ID: {run.id}")
    y -= 6 * mm
    completed = run.completed_at or datetime.now(UTC)
    c.drawString(25 * mm, y, f"Completed: {completed:%Y-%m-%d %H:%M UTC}")
    y -= 12 * mm

    c.setFont("Helvetica-Bold", 11)
    c.drawString(25 * mm, y, "Totals")
    y -= 7 * mm
    c.setFont("Helvetica", 10)
    for label, value in totals:
        c.drawString(28 * mm, y, label)
        c.drawRightString(width - 30 * mm, y, str(value))
        y -= 6 * mm

    y -= 6 * mm
    c.setFont("Helvetica-Bold", 11)
    c.drawString(25 * mm, y, f"Violation detail ({len(violations)})")
    y -= 6 * mm
    c.setFont("Helvetica", 8.5)
    for v in violations[:40]:
        line = (
            f"{v.employee_code}  {(v.employee_name or '')[:24]:24}  "
            f"{v.email_status.value:16}  {v.email_to or '—'}"
        )
        c.drawString(25 * mm, y, line[:110])
        y -= 4.5 * mm
        if y < 20 * mm:
            c.showPage()
            y = height - 25 * mm
            c.setFont("Helvetica", 8.5)
    c.showPage()
    c.save()
    return buf.getvalue()


async def generate_and_store_reports(db: AsyncSession, run: AnalysisRun) -> dict:
    violations = await _run_rows(db, run)
    skips = run.skipped_rows or []

    storage = get_storage()
    made: dict[str, str] = {}

    xlsx_bytes = build_excel_3tab(run, violations, skips)
    key_xlsx = build_key(
        manager_id=str(run.manager_id),
        kind="GETS",
        batch_id=str(run.batch_id),
        filename=f"analysis-{str(run.id)[:8]}-3tab.xlsx",
    ).replace("/gets/", "/reports/")  # keep reports out of upload trees
    storage.put(key_xlsx, io.BytesIO(xlsx_bytes), len(xlsx_bytes))
    db.add(
        GeneratedReport(
            run_id=run.id, kind=ReportKind.EXCEL_3TAB,
            storage_key=key_xlsx, emailed_to=None,
        )
    )
    made["excel_3tab"] = key_xlsx

    pdf_bytes = build_summary_pdf(run, violations)
    key_pdf = build_key(
        manager_id=str(run.manager_id),
        kind="GETS",
        batch_id=str(run.batch_id),
        filename=f"analysis-{str(run.id)[:8]}-summary.pdf",
    ).replace("/gets/", "/reports/")
    storage.put(key_pdf, io.BytesIO(pdf_bytes), len(pdf_bytes))
    db.add(
        GeneratedReport(
            run_id=run.id, kind=ReportKind.SUMMARY_PDF,
            storage_key=key_pdf, emailed_to=None,
        )
    )
    made["summary_pdf"] = key_pdf

    await db.commit()

    # Email both artifacts to the uploading manager (dry-run aware).
    manager = await db.get(User, run.manager_id)
    await email_reports_to_manager(manager, made)
    return made


async def email_reports_to_manager(manager: User | None, keys: dict[str, str]) -> None:
    if manager is None or not manager.email:
        return
    from app.core.config import settings as cfg

    log.info(
        "report_delivery",
        mode=cfg.outbound_email_mode,
        to=manager.email,
        attachments=sorted(keys.values()),
    )
